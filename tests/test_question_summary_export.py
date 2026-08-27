"""Tests for the readable Question Bank summary export."""
from __future__ import annotations

import io
import os
import sys
import unittest
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from unittest.mock import Mock, patch

from openpyxl import load_workbook
from flask import Flask
from flask_login import LoginManager

PROJECT_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROJECT_ROOT))

os.environ.setdefault('FLASK_ENV', 'testing')
os.environ['DATABASE_URL'] = 'sqlite:///:memory:'

from app.models import db
from app.models.admin import Admin
from app.models.course import College
from app.routes import questions as question_routes
from app.routes.questions import QuestionBankService


def build_summary_service():
    service = QuestionBankService('/tmp')
    service.questions = {
        'Q_SEL': {
            'id': 'Q_SEL',
            'type_id': 'T_SEL',
            'type': 'Selection',
            'text': 'The course challenged me to think critically.',
            'detail': 'Shown near the end of the survey.',
            'block_title': 'Course Experience',
        },
        'Q_COM': {
            'id': 'Q_COM',
            'type_id': 'T_COM',
            'type': 'Comment',
            'text': 'What helped your learning most?',
            'detail': '',
            'block_title': '',
        },
    }
    service.question_types = {
        'T_SEL': {
            'id': 'T_SEL',
            'name': 'Agreement',
            'type': 'Selection',
            'options': ['Strongly disagree', 'Disagree', 'Agree', 'Strongly agree'],
        },
        'T_COM': {
            'id': 'T_COM',
            'name': 'Open comment',
            'type': 'Comment',
            'options': [],
        },
    }
    service.question_mapping = defaultdict(dict, {
        'DEPARTMENT': {'AS-DEPT': {'CRS_SEL_01': 'Q_SEL'}},
        'COURSE': {'EGR 101': {'CRS_INS_COM_01': 'Q_COM'}},
        'SECTION': {'AS111-001': {'CRS_COM_01': 'Q_COM'}},
    })
    service._unit_to_college = {
        'AS-DEPT': 'AS',
        'AS111-001': 'AS',
        'EGR 101': 'EN',
    }
    service._unit_metadata = {
        ('DEPARTMENT', 'AS-DEPT'): {
            'college_code': 'AS',
            'college_name': 'Arts and Sciences',
            'department_name': 'Arts and Sciences',
            'course': 'SHOULD NOT APPEAR',
            'section': 'SHOULD NOT APPEAR',
            'section_title': 'SHOULD NOT APPEAR',
        },
        ('SECTION', 'AS111-001'): {
            'college_code': 'AS',
            'college_name': 'Arts and Sciences',
            'department_name': 'Arts and Sciences',
            'course': 'A&S 111',
            'section': 'AS111-001',
            'section_title': 'Culture and Society',
        },
        ('COURSE', 'EGR 101'): {
            'college_code': 'EN',
            'college_name': 'Engineering',
            'department_name': 'Engineering',
            'course': 'EGR 101',
            'section': 'SHOULD NOT APPEAR',
            'section_title': 'SHOULD NOT APPEAR',
        },
    }
    service.hierarchy = {
        'Arts and Sciences': {'id': 'AS', 'type': 'college', 'children': {}},
        'Engineering': {'id': 'EN', 'type': 'college', 'children': {}},
    }
    return service


class QuestionSummaryWorkbookTests(unittest.TestCase):
    def test_super_admin_workbook_has_one_sheet_per_college(self):
        output, summary = build_summary_service().export_question_bank_summary(
            exported_at=datetime(2026, 8, 19, 10, 30)
        )
        workbook = load_workbook(output)

        self.assertEqual(summary['row_count'], 3)
        self.assertEqual(summary['sheet_count'], 2)
        self.assertEqual(len(workbook.sheetnames), 2)
        self.assertTrue(any(name.startswith('AS - ') for name in workbook.sheetnames))
        self.assertTrue(any(name.startswith('EN - ') for name in workbook.sheetnames))

        as_sheet = next(sheet for sheet in workbook.worksheets if sheet.title.startswith('AS - '))
        self.assertEqual(as_sheet['A4'].value, 'Unit Level')
        self.assertEqual(as_sheet['G4'].value, 'Question')
        self.assertEqual(as_sheet.freeze_panes, 'A5')
        self.assertFalse(as_sheet.sheet_view.showGridLines)
        self.assertEqual(as_sheet.page_setup.orientation, 'landscape')
        self.assertEqual(len(as_sheet.tables), 1)

        department_row = next(
            row for row in as_sheet.iter_rows(min_row=5, values_only=True)
            if row[0] == 'Department'
        )
        self.assertEqual(department_row[1], 'Arts and Sciences')
        self.assertIsNone(department_row[2])
        self.assertIsNone(department_row[3])
        self.assertIn('Strongly disagree', department_row[7])

    def test_college_workbook_is_single_sheet_and_fails_closed(self):
        output, summary = build_summary_service().export_question_bank_summary(
            college_code='AS',
            exported_at=datetime(2026, 8, 19, 10, 30),
        )
        workbook = load_workbook(output)

        self.assertEqual(summary['row_count'], 2)
        self.assertEqual(len(workbook.sheetnames), 1)
        values = [cell.value for row in workbook.active.iter_rows() for cell in row]
        self.assertNotIn('EGR 101', values)
        self.assertNotIn('What helped your learning most?', [
            row[6].value for row in workbook.active.iter_rows(min_row=5)
            if len(row) > 6 and row[0].value == 'Course'
        ])


class QuestionSummaryRouteTests(unittest.TestCase):
    def setUp(self):
        self.app = Flask(__name__)
        self.app.config.update(
            TESTING=True,
            SECRET_KEY='question-summary-tests',
            SQLALCHEMY_DATABASE_URI='sqlite:///:memory:',
            SQLALCHEMY_TRACK_MODIFICATIONS=False,
        )
        db.init_app(self.app)
        login_manager = LoginManager()
        login_manager.init_app(self.app)

        @login_manager.user_loader
        def load_user(user_id):
            return db.session.get(Admin, int(user_id))

        self.app.register_blueprint(question_routes.questions_bp, url_prefix='/questions')
        self.context = self.app.app_context()
        self.context.push()
        self.previous_service = question_routes._qb_service

        db.create_all()

        db.session.add_all([
            College(code='AS', name='Arts and Sciences', qb_enabled=True),
            College(code='EN', name='Engineering', qb_enabled=True),
        ])
        self.college_admin = Admin(
            linkblue='college-admin', first_name='College', last_name='Admin',
            role='college_admin', college_code='AS', has_qb_access=True,
        )
        self.department_admin = Admin(
            linkblue='department-admin', first_name='Department', last_name='Admin',
            role='dept_admin', college_code='AS', has_qb_access=True,
        )
        self.super_admin = Admin(
            linkblue='super-admin', first_name='Super', last_name='Admin',
            role='super_admin', has_qb_access=True,
        )
        db.session.add_all([self.college_admin, self.department_admin, self.super_admin])
        db.session.commit()

        self.fake_service = Mock()
        self.fake_service.export_question_bank_summary.return_value = (
            io.BytesIO(b'workbook-bytes'),
            {'row_count': 2, 'sheet_count': 1, 'sheet_names': ['AS']},
        )
        question_routes._qb_service = self.fake_service
        self.client = self.app.test_client()

    def tearDown(self):
        question_routes._qb_service = self.previous_service
        db.session.remove()
        db.drop_all()
        self.context.pop()

    def login(self, admin):
        with self.client.session_transaction() as session:
            session['_user_id'] = str(admin.id)
            session['_fresh'] = True

    @patch.object(question_routes, 'log_audit')
    @patch.object(question_routes.QBAuditLog, 'log_action')
    def test_college_admin_export_is_scoped(self, _log_db, _log_file):
        self.login(self.college_admin)
        response = self.client.get('/questions/export/summary')

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data, b'workbook-bytes')
        self.assertIn('QB_summary_AS_', response.headers['Content-Disposition'])
        self.fake_service.load_courses.assert_called_once_with(
            admin_scope={'college': 'AS', 'department': None}
        )
        self.fake_service.export_question_bank_summary.assert_called_once_with(college_code='AS')

    def test_department_admin_cannot_export_summary(self):
        self.login(self.department_admin)
        response = self.client.get('/questions/export/summary')

        self.assertEqual(response.status_code, 302)
        self.fake_service.export_question_bank_summary.assert_not_called()

    @patch.object(question_routes, 'log_audit')
    @patch.object(question_routes.QBAuditLog, 'log_action')
    def test_super_admin_export_includes_all_colleges(self, _log_db, _log_file):
        self.login(self.super_admin)
        response = self.client.get('/questions/export/summary')

        self.assertEqual(response.status_code, 200)
        self.assertIn('QB_summary_all_colleges_', response.headers['Content-Disposition'])
        self.fake_service.load_courses.assert_called_once_with(admin_scope=None)
        self.fake_service.export_question_bank_summary.assert_called_once_with(college_code=None)


if __name__ == '__main__':
    unittest.main()
